"""Discovery documents: parsing, grounding, upload → extract → import, retention."""

import io
import json
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from cryptography.fernet import Fernet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.config import get_settings
from app.models.company import Company
from app.models.contact import Contact
from app.models.discovery_document import DiscoveryDocument
from app.models.org_llm_key import OrgLlmKey
from app.models.tenant import Organization
from app.models.user import UserRole
from app.org_context import APPREND_ORG_ID
from app.routers import discovery_documents as router_mod
from app.services import crypto, discovery_documents, document_extract, org_gemini
from app.services.document_extract import DocumentRejected, TableMapping, TextExtraction

CSV = (
    "Company,Website,Contact,Email,Title\n"
    "Acme Stamping,https://www.acme-stamping.com,Jane Doe,jane@acme-stamping.com,COO\n"
    "Acme Stamping,acme-stamping.com,Raj Patel,raj@acme-stamping.com,Plant Manager\n"
    "Beta Molding,,Lee Wong,lee.wong@gmail.com,Owner\n"
    ",,,,\n"
).encode()

CSV_MAPPING = {
    "has_header": True,
    "mapping": {"company_name": 0, "website": 1, "full_name": 2, "email": 3, "title": 4},
}


@pytest.fixture(autouse=True)
def _crypto_key():
    settings = get_settings()
    prev = settings.token_encryption_key
    settings.token_encryption_key = Fernet.generate_key().decode()
    crypto._fernet.cache_clear()
    yield
    settings.token_encryption_key = prev
    crypto._fernet.cache_clear()


@pytest.fixture(autouse=True)
def _reset_org_breaker():
    org_gemini.reset_breaker()
    yield


@pytest.fixture
def fake_gemini(monkeypatch):
    """Answer table prompts with CSV_MAPPING; record every prompt."""
    calls = []

    async def fake_chat(messages, *, providers, **kw):
        prompt = messages[0]["content"]
        calls.append(prompt)
        return json.dumps(CSV_MAPPING), providers[0].name, providers[0].models[0]

    monkeypatch.setattr(org_gemini, "chat_completion", fake_chat)
    return calls


@pytest.fixture
def bg_sessions(db_engine, monkeypatch):
    monkeypatch.setattr(
        router_mod, "session_factory",
        async_sessionmaker(db_engine, class_=AsyncSession, expire_on_commit=False),
    )


async def _save_key(db_session, org_id=APPREND_ORG_ID):
    db_session.add(OrgLlmKey(
        org_id=org_id, provider="gemini",
        key_encrypted=crypto.encrypt("test-gemini-key-123"), key_hint="••••",
    ))
    await db_session.commit()


# ── classify / parse ─────────────────────────────────────────────────────────

def test_classify_checks_bytes_not_just_extension():
    assert document_extract.classify("leads.csv", CSV) == "csv"
    with pytest.raises(DocumentRejected):
        document_extract.classify("fake.pdf", b"hello")
    with pytest.raises(DocumentRejected):
        document_extract.classify("leads.exe", b"MZ")
    with pytest.raises(DocumentRejected):
        document_extract.classify("leads.csv", b"a,b\x00c")
    with pytest.raises(DocumentRejected):
        document_extract.classify("sheet.xlsx", b"PK\x03\x04garbage")
    with pytest.raises(DocumentRejected):
        document_extract.classify("big.txt", b"x" * (document_extract.MAX_FILE_BYTES + 1))


def test_table_mapping_is_applied_to_every_row_and_merges_companies():
    doc = document_extract.parse_document("csv", CSV)
    candidates, stats = document_extract.candidates_from_table(doc, TableMapping(**CSV_MAPPING))
    assert [c["name"] for c in candidates] == ["Acme Stamping", "Beta Molding"]
    acme, beta = candidates
    assert acme["domain"] == "acme-stamping.com"
    assert [(p["first_name"], p["last_name"], p["email"]) for p in acme["contacts"]] == [
        ("Jane", "Doe", "jane@acme-stamping.com"),
        ("Raj", "Patel", "raj@acme-stamping.com"),
    ]
    assert acme["contacts"][0]["source_ref"] == "row 2"
    # A free-mail address is kept as the contact's email but never becomes the domain.
    assert beta["domain"] is None
    assert beta["contacts"][0]["email"] == "lee.wong@gmail.com"


def test_mapping_to_columns_that_do_not_exist_is_rejected():
    doc = document_extract.parse_document("csv", CSV)
    with pytest.raises(DocumentRejected):
        document_extract.candidates_from_table(
            doc, TableMapping(has_header=True, mapping={"email": 42, "bogus": 0})
        )


def test_xlsx_and_docx_parse():
    from docx import Document
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Company", "Phone"])
    wb.active.append(["Gamma Tool", 5551234567.0])
    buf = io.BytesIO()
    wb.save(buf)
    doc = document_extract.parse_document("xlsx", buf.getvalue())
    assert doc.kind == "table" and doc.rows == [["Gamma Tool", "5551234567"]]

    d = Document()
    d.add_paragraph("Met Priya Shah (priya@delta-die.com), VP Ops at Delta Die Works.")
    buf = io.BytesIO()
    d.save(buf)
    assert document_extract.classify("notes.docx", buf.getvalue()) == "docx"
    doc = document_extract.parse_document("docx", buf.getvalue())
    assert doc.kind == "text" and "priya@delta-die.com" in doc.chunks[0][1]


def test_textless_pdf_is_rejected_as_scanned():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(DocumentRejected, match="Scanned"):
        document_extract.parse_document("pdf", buf.getvalue())


def test_text_extraction_drops_identifiers_not_in_the_text():
    text = "[part 1]\nMet Priya Shah, VP Ops at Delta Die Works. Call 555-201-3344."
    b = document_extract._Builder()
    answer = TextExtraction.model_validate({"companies": [{
        "name": "Delta Die Works",
        "domain": "deltadie.com",              # not in the text
        "phone": "(555) 201-3344",              # same digits as the text
        "confidence": 0.9,
        "contacts": [{"first_name": "Priya", "last_name": "Shah", "title": "VP Ops",
                      "email": "priya@deltadie.com"}],  # invented
    }]})
    fields, contacts = document_extract._grounded(answer.companies[0], text, b)
    assert fields["domain"] is None
    assert fields["phone"] == "(555) 201-3344"
    assert contacts[0]["email"] is None
    assert contacts[0]["first_name"] == "Priya"
    assert b.dropped_fields == 2


# ── API flow ─────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_upload_without_gemini_key_is_409_and_stores_nothing(client, db_session):
    resp = await client.post(
        "/api/v1/discovery/documents", files={"file": ("leads.csv", CSV, "text/csv")}
    )
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "gemini_key_required"
    assert (await db_session.execute(select(DiscoveryDocument))).first() is None
    status = await client.get("/api/v1/discovery/documents/ai-status")
    assert status.json() == {"gemini_key_configured": False}


@pytest.mark.asyncio
async def test_upload_extract_and_import(client, db_session, fake_gemini, bg_sessions):
    await _save_key(db_session)
    resp = await client.post(
        "/api/v1/discovery/documents", files={"file": ("leads.csv", CSV, "text/csv")}
    )
    assert resp.status_code == 202, resp.text
    doc_id = resp.json()["id"]

    detail = (await client.get(f"/api/v1/discovery/documents/{doc_id}")).json()
    assert detail["status"] == "ready", detail
    assert detail["candidate_count"] == 2
    # Gemini only ever saw the header + sample rows, never a whole table dump.
    assert len(fake_gemini) == 1
    acme = detail["candidates"][0]
    assert acme["match"] is None

    imp = await client.post(
        f"/api/v1/discovery/documents/{doc_id}/import", json={"candidate_ids": [acme["id"]]}
    )
    assert imp.status_code == 200, imp.text
    body = imp.json()
    assert body["companies_created"] == 1
    assert body["contacts_created"] == 2
    assert len(body["contact_ids"]) == 2

    company = (await db_session.execute(
        select(Company).where(Company.domain == "acme-stamping.com")
    )).scalar_one()
    assert company.source == "document"
    emails = sorted((await db_session.execute(
        select(Contact.email).where(Contact.company_id == company.id)
    )).scalars().all())
    assert emails == ["jane@acme-stamping.com", "raj@acme-stamping.com"]

    after = (await client.get(f"/api/v1/discovery/documents/{doc_id}")).json()
    assert after["candidates"][0]["imported_company_id"] == str(company.id)
    assert after["candidates"][0]["contacts"][0]["imported_contact_id"]

    # Re-uploading the same file now flags Acme as already in Prospect.
    resp2 = await client.post(
        "/api/v1/discovery/documents", files={"file": ("leads.csv", CSV, "text/csv")}
    )
    again = (await client.get(f"/api/v1/discovery/documents/{resp2.json()['id']}")).json()
    assert again["candidates"][0]["match"]["matched_company_id"] == str(company.id)


@pytest.mark.asyncio
async def test_import_only_chosen_contacts(client, db_session, fake_gemini, bg_sessions):
    await _save_key(db_session)
    resp = await client.post(
        "/api/v1/discovery/documents", files={"file": ("leads.csv", CSV, "text/csv")}
    )
    doc_id = resp.json()["id"]
    acme = (await client.get(f"/api/v1/discovery/documents/{doc_id}")).json()["candidates"][0]
    first, second = acme["contacts"][0], acme["contacts"][1]

    imp = await client.post(
        f"/api/v1/discovery/documents/{doc_id}/import",
        json={"candidate_ids": [acme["id"]], "contact_ids": [first["id"]]},
    )
    assert imp.status_code == 200, imp.text
    assert imp.json()["contacts_created"] == 1
    assert len(imp.json()["contact_ids"]) == 1

    company = (await db_session.execute(
        select(Company).where(Company.domain == "acme-stamping.com")
    )).scalar_one()
    emails = (await db_session.execute(
        select(Contact.email).where(Contact.company_id == company.id)
    )).scalars().all()
    assert emails == [first["email"]]
    after = (await client.get(f"/api/v1/discovery/documents/{doc_id}")).json()["candidates"][0]
    assert after["contacts"][0]["imported_contact_id"]
    assert not after["contacts"][1].get("imported_contact_id")

    # Coming back for the second person merges into the same company.
    again = await client.post(
        f"/api/v1/discovery/documents/{doc_id}/import",
        json={"candidate_ids": [acme["id"]], "contact_ids": [second["id"]]},
    )
    assert again.json()["companies_created"] == 0
    assert again.json()["contacts_created"] == 1
    final = (await client.get(f"/api/v1/discovery/documents/{doc_id}")).json()["candidates"][0]
    assert all(c.get("imported_contact_id") for c in final["contacts"])


@pytest.mark.asyncio
async def test_bad_file_is_422(client, db_session):
    await _save_key(db_session)
    resp = await client.post(
        "/api/v1/discovery/documents", files={"file": ("x.pdf", b"not a pdf", "application/pdf")}
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_documents_are_invisible_to_other_orgs(
    client, db_session, fake_gemini, bg_sessions, make_user, auth_headers
):
    await _save_key(db_session)
    doc_id = (await client.post(
        "/api/v1/discovery/documents", files={"file": ("leads.csv", CSV, "text/csv")}
    )).json()["id"]

    other = Organization(id=uuid.uuid4(), name="Other", slug="other-org",
                         is_active=True, billing_plan="pro")
    db_session.add(other)
    await db_session.commit()
    outsider = await make_user("outsider@other.test", role=UserRole.ORG_ADMIN.value,
                               org_id=other.id)
    h = auth_headers(outsider)
    assert (await client.get(f"/api/v1/discovery/documents/{doc_id}", headers=h)).status_code == 404
    assert (await client.get("/api/v1/discovery/documents", headers=h)).json() == []
    imp = await client.post(f"/api/v1/discovery/documents/{doc_id}/import",
                            json={"candidate_ids": ["c1"]}, headers=h)
    assert imp.status_code == 404
    # And the other org has no key, so it can't upload on ours.
    up = await client.post("/api/v1/discovery/documents",
                           files={"file": ("leads.csv", CSV, "text/csv")}, headers=h)
    assert up.status_code == 409


@pytest.mark.asyncio
async def test_gemini_failure_marks_document_failed(client, db_session, bg_sessions, monkeypatch):
    await _save_key(db_session)

    async def broken(messages, *, providers, **kw):
        from app.services.llm_providers import LLMChainExhausted
        raise LLMChainExhausted("HTTP 400 API key not valid")

    monkeypatch.setattr(org_gemini, "chat_completion", broken)
    doc_id = (await client.post(
        "/api/v1/discovery/documents", files={"file": ("leads.csv", CSV, "text/csv")}
    )).json()["id"]
    detail = (await client.get(f"/api/v1/discovery/documents/{doc_id}")).json()
    assert detail["status"] == "failed"
    assert "Gemini API key" in detail["error"]


# ── retention / staleness ────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_purge_drops_bytes_after_30_days_and_keeps_results(db_session):
    old = DiscoveryDocument(
        org_id=APPREND_ORG_ID, filename="old.csv", file_type="csv", size_bytes=3,
        content=b"a,b", status="ready", extracted={"candidates": [{"id": "c1"}]},
        created_at=datetime.now(timezone.utc) - timedelta(days=31),
    )
    new = DiscoveryDocument(
        org_id=APPREND_ORG_ID, filename="new.csv", file_type="csv", size_bytes=3,
        content=b"a,b", status="ready",
    )
    db_session.add_all([old, new])
    await db_session.commit()

    assert await discovery_documents.purge_expired_files(db_session) == 1
    await db_session.commit()
    rows = {
        name: (content, extracted, purged)
        for name, content, extracted, purged in (await db_session.execute(
            select(DiscoveryDocument.filename, DiscoveryDocument.content,
                   DiscoveryDocument.extracted, DiscoveryDocument.content_purged_at)
        )).all()
    }
    assert rows["old.csv"][0] is None and rows["old.csv"][2] is not None
    assert rows["old.csv"][1] == {"candidates": [{"id": "c1"}]}
    assert rows["new.csv"][0] == b"a,b"


@pytest.mark.asyncio
async def test_stuck_parsing_document_becomes_retryable_failure(db_session):
    doc = DiscoveryDocument(
        org_id=APPREND_ORG_ID, filename="x.csv", file_type="csv", size_bytes=1,
        content=b"a", status="parsing",
        created_at=datetime.now(timezone.utc) - timedelta(hours=1),
    )
    db_session.add(doc)
    await db_session.commit()
    await discovery_documents.mark_stale(db_session, doc)
    assert doc.status == "failed"
